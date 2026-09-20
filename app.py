from __future__ import annotations

import csv
import io
import os
import sqlite3
from pathlib import Path
from typing import Optional

from fastapi import FastAPI, HTTPException, Query, Request
from fastapi.responses import HTMLResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from openpyxl import Workbook
from pydantic import BaseModel, Field

BASE_DIR = Path(__file__).resolve().parent
DB_PATH = Path(os.environ.get("SIEMPRO_DB_PATH", BASE_DIR / "siempro.db"))

app = FastAPI(title="SIEMPro", version="1.0.0")
app.mount("/static", StaticFiles(directory=str(BASE_DIR / "static")), name="static")
templates = Jinja2Templates(directory=str(BASE_DIR / "templates"))


class EventCreate(BaseModel):
    name: str = Field(..., min_length=2)
    date: str = Field(..., min_length=4)
    location: str = Field(..., min_length=2)
    description: Optional[str] = ""


class CompetitionCreate(BaseModel):
    event_id: int
    name: str = Field(..., min_length=2)
    time: str = Field(..., min_length=2)
    venue: str = Field(..., min_length=2)
    capacity: int = Field(..., gt=0)
    coordinator_name: str = Field(..., min_length=2)
    description: Optional[str] = ""
    rules: Optional[str] = ""


class CompetitionMessageCreate(BaseModel):
    event_id: int
    competition_id: int
    sender: str = Field(..., min_length=2)
    sender_name: str = Field(..., min_length=2)
    message: str = Field(..., min_length=2)


class ParticipantCreate(BaseModel):
    event_id: int
    competition_id: int
    name: str = Field(..., min_length=2)
    email: str = Field(..., min_length=4)
    phone: str = Field(..., min_length=5)
    department: str = Field(..., min_length=2)
    year: str = Field(..., min_length=1)
    registration_no: str = Field(..., min_length=3)


class VolunteerCreate(BaseModel):
    event_id: int
    competition_id: int
    name: str = Field(..., min_length=2)
    email: str = Field(..., min_length=4)
    phone: str = Field(..., min_length=5)
    preferred_role: str = Field(..., min_length=2)
    skills: str = Field(..., min_length=2)


class LoginRequest(BaseModel):
    username: str = Field(..., min_length=2)
    password: str = Field(..., min_length=2)


def get_connection() -> sqlite3.Connection:
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def init_db() -> None:
    with get_connection() as conn:
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS events (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL,
                date TEXT NOT NULL,
                location TEXT NOT NULL,
                description TEXT DEFAULT ''
            )
            """
        )
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS users (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                username TEXT NOT NULL UNIQUE,
                password TEXT NOT NULL,
                role TEXT NOT NULL DEFAULT 'admin'
            )
            """
        )
        conn.execute(
            "INSERT OR IGNORE INTO users (username, password, role) VALUES (?, ?, ?)",
            ("admin", "admin123", "admin"),
        )
        conn.execute(
            "INSERT OR IGNORE INTO users (username, password, role) VALUES (?, ?, ?)",
            ("coordinator", "coord123", "coordinator"),
        )

        default_roles = {"admin": "admin123", "coordinator": "coord123"}
        existing_users = {
            row["username"]: row["password"] for row in conn.execute("SELECT username, password FROM users").fetchall()
        }
        for username, password in default_roles.items():
            if username not in existing_users:
                conn.execute(
                    "INSERT INTO users (username, password, role) VALUES (?, ?, ?)",
                    (username, password, username if username == "admin" else "coordinator"),
                )

        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS competitions (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                event_id INTEGER NOT NULL,
                name TEXT NOT NULL,
                time TEXT NOT NULL,
                venue TEXT NOT NULL,
                capacity INTEGER NOT NULL,
                coordinator_name TEXT NOT NULL,
                description TEXT DEFAULT '',
                rules TEXT DEFAULT '',
                FOREIGN KEY (event_id) REFERENCES events(id) ON DELETE CASCADE
            )
            """
        )
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS participants (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                event_id INTEGER NOT NULL,
                competition_id INTEGER NOT NULL,
                name TEXT NOT NULL,
                email TEXT NOT NULL,
                phone TEXT NOT NULL,
                department TEXT NOT NULL,
                year TEXT NOT NULL,
                registration_no TEXT NOT NULL,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY (event_id) REFERENCES events(id) ON DELETE CASCADE,
                FOREIGN KEY (competition_id) REFERENCES competitions(id) ON DELETE CASCADE
            )
            """
        )
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS volunteers (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                event_id INTEGER NOT NULL,
                competition_id INTEGER NOT NULL,
                name TEXT NOT NULL,
                email TEXT NOT NULL,
                phone TEXT NOT NULL,
                preferred_role TEXT NOT NULL,
                skills TEXT NOT NULL,
                assignment TEXT DEFAULT 'Pending',
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY (event_id) REFERENCES events(id) ON DELETE CASCADE,
                FOREIGN KEY (competition_id) REFERENCES competitions(id) ON DELETE CASCADE
            )
            """
        )
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS competition_messages (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                event_id INTEGER NOT NULL,
                competition_id INTEGER NOT NULL,
                sender TEXT NOT NULL,
                sender_name TEXT NOT NULL,
                message TEXT NOT NULL,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY (event_id) REFERENCES events(id) ON DELETE CASCADE,
                FOREIGN KEY (competition_id) REFERENCES competitions(id) ON DELETE CASCADE
            )
            """
        )

        competition_columns = [
            row[1] for row in conn.execute("PRAGMA table_info(competitions)").fetchall()
        ]
        if "rules" not in competition_columns:
            conn.execute("ALTER TABLE competitions ADD COLUMN rules TEXT DEFAULT ''")

        conn.commit()


@app.on_event("startup")
def startup_event() -> None:
    init_db()


@app.get("/", response_class=HTMLResponse)
async def index(request: Request):
    return templates.TemplateResponse("index.html", {"request": request})


@app.post("/api/login")
async def login(payload: LoginRequest):
    with get_connection() as conn:
        row = conn.execute(
            "SELECT username, role FROM users WHERE username = ? AND password = ?",
            (payload.username.strip(), payload.password),
        ).fetchone()
    if not row:
        raise HTTPException(status_code=401, detail="Invalid username or password.")
    return {"username": row["username"], "role": row["role"], "message": "Login successful."}


@app.get("/api/health")
async def health_check():
    return {"status": "ok", "service": "SIEMPro"}


@app.get("/api/events")
async def list_events():
    with get_connection() as conn:
        rows = conn.execute(
            "SELECT * FROM events ORDER BY date DESC, id DESC"
        ).fetchall()
    return [dict(row) for row in rows]


@app.post("/api/events")
async def create_event(payload: EventCreate):
    with get_connection() as conn:
        cursor = conn.execute(
            "INSERT INTO events (name, date, location, description) VALUES (?, ?, ?, ?)",
            (payload.name.strip(), payload.date, payload.location.strip(), payload.description or ""),
        )
        conn.commit()
        event_id = cursor.lastrowid
    return {"id": event_id, "message": "Event created successfully."}


@app.delete("/api/events/{event_id}")
async def delete_event(event_id: int):
    with get_connection() as conn:
        cursor = conn.execute("DELETE FROM events WHERE id = ?", (event_id,))
        conn.commit()
    if cursor.rowcount == 0:
        raise HTTPException(status_code=404, detail="Event not found.")
    return {"message": "Event deleted successfully."}


@app.get("/api/competitions")
async def list_competitions(event_id: Optional[int] = Query(None)):
    with get_connection() as conn:
        if event_id is not None:
            rows = conn.execute(
                "SELECT * FROM competitions WHERE event_id = ? ORDER BY time ASC, id DESC",
                (event_id,),
            ).fetchall()
        else:
            rows = conn.execute(
                "SELECT * FROM competitions ORDER BY event_id ASC, time ASC"
            ).fetchall()
    return [dict(row) for row in rows]


@app.post("/api/competitions")
async def create_competition(payload: CompetitionCreate):
    with get_connection() as conn:
        event_exists = conn.execute("SELECT 1 FROM events WHERE id = ?", (payload.event_id,)).fetchone()
        if not event_exists:
            raise HTTPException(status_code=404, detail="Event not found.")
        cursor = conn.execute(
            """
            INSERT INTO competitions (event_id, name, time, venue, capacity, coordinator_name, description, rules)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                payload.event_id,
                payload.name.strip(),
                payload.time,
                payload.venue.strip(),
                payload.capacity,
                payload.coordinator_name.strip(),
                payload.description or "",
                payload.rules or "",
            ),
        )
        conn.commit()
        competition_id = cursor.lastrowid
    return {"id": competition_id, "message": "Competition created successfully."}


@app.patch("/api/competitions/{competition_id}")
async def update_competition(competition_id: int, payload: dict):
    with get_connection() as conn:
        existing = conn.execute("SELECT * FROM competitions WHERE id = ?", (competition_id,)).fetchone()
        if not existing:
            raise HTTPException(status_code=404, detail="Competition not found.")

        updates = {}
        if "name" in payload and payload["name"]:
            updates["name"] = payload["name"].strip()
        if "capacity" in payload:
            updates["capacity"] = int(payload["capacity"])
        if "rules" in payload:
            updates["rules"] = payload["rules"].strip() if payload["rules"] else ""
        if "description" in payload:
            updates["description"] = payload["description"].strip() if payload["description"] else ""

        if not updates:
            raise HTTPException(status_code=400, detail="No valid update fields provided.")

        fields = []
        values = []
        for key, value in updates.items():
            fields.append(f"{key} = ?")
            values.append(value)
        values.append(competition_id)

        conn.execute(f"UPDATE competitions SET {', '.join(fields)} WHERE id = ?", tuple(values))
        conn.commit()
        row = conn.execute("SELECT * FROM competitions WHERE id = ?", (competition_id,)).fetchone()
    return dict(row)


@app.get("/api/competition-rules")
async def get_competition_rules(event_id: int = Query(...), competition_id: int = Query(...)):
    with get_connection() as conn:
        row = conn.execute(
            "SELECT rules, name FROM competitions WHERE id = ? AND event_id = ?",
            (competition_id, event_id),
        ).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="Competition not found.")
    return {"competition_name": row["name"], "rules": row["rules"] or "No rules added yet."}


@app.get("/api/competition-messages")
async def list_competition_messages(event_id: int = Query(...), competition_id: int = Query(...)):
    with get_connection() as conn:
        rows = conn.execute(
            """
            SELECT id, event_id, competition_id, sender, sender_name, message, created_at
            FROM competition_messages
            WHERE event_id = ? AND competition_id = ?
            ORDER BY created_at ASC
            """,
            (event_id, competition_id),
        ).fetchall()
    return [dict(row) for row in rows]


@app.post("/api/competition-messages")
async def create_competition_message(payload: CompetitionMessageCreate):
    with get_connection() as conn:
        competition = conn.execute(
            "SELECT id FROM competitions WHERE id = ? AND event_id = ?",
            (payload.competition_id, payload.event_id),
        ).fetchone()
        if not competition:
            raise HTTPException(status_code=404, detail="Competition not found for the selected event.")

        cursor = conn.execute(
            """
            INSERT INTO competition_messages (event_id, competition_id, sender, sender_name, message)
            VALUES (?, ?, ?, ?, ?)
            """,
            (
                payload.event_id,
                payload.competition_id,
                payload.sender.strip().lower(),
                payload.sender_name.strip(),
                payload.message.strip(),
            ),
        )
        conn.commit()
        message_id = cursor.lastrowid
    return {"id": message_id, "message": "Message saved successfully."}


@app.get("/api/participants")
async def list_participants(competition_id: Optional[int] = Query(None), event_id: Optional[int] = Query(None)):
    with get_connection() as conn:
        if competition_id is not None:
            rows = conn.execute(
                "SELECT * FROM participants WHERE competition_id = ? ORDER BY created_at DESC",
                (competition_id,),
            ).fetchall()
        elif event_id is not None:
            rows = conn.execute(
                "SELECT * FROM participants WHERE event_id = ? ORDER BY created_at DESC",
                (event_id,),
            ).fetchall()
        else:
            rows = conn.execute(
                "SELECT * FROM participants ORDER BY created_at DESC"
            ).fetchall()
    return [dict(row) for row in rows]


@app.get("/api/reports/participants/csv")
async def export_participants_csv(event_id: Optional[int] = Query(None), competition_id: Optional[int] = Query(None)):
    with get_connection() as conn:
        if competition_id is not None:
            rows = conn.execute(
                "SELECT p.name, p.email, p.phone, p.department, p.year, p.registration_no, e.name AS event_name, c.name AS competition_name FROM participants p JOIN events e ON e.id = p.event_id JOIN competitions c ON c.id = p.competition_id WHERE p.competition_id = ? ORDER BY p.created_at DESC",
                (competition_id,),
            ).fetchall()
        elif event_id is not None:
            rows = conn.execute(
                "SELECT p.name, p.email, p.phone, p.department, p.year, p.registration_no, e.name AS event_name, c.name AS competition_name FROM participants p JOIN events e ON e.id = p.event_id JOIN competitions c ON c.id = p.competition_id WHERE p.event_id = ? ORDER BY p.created_at DESC",
                (event_id,),
            ).fetchall()
        else:
            rows = conn.execute(
                "SELECT p.name, p.email, p.phone, p.department, p.year, p.registration_no, e.name AS event_name, c.name AS competition_name FROM participants p JOIN events e ON e.id = p.event_id JOIN competitions c ON c.id = p.competition_id ORDER BY p.created_at DESC"
            ).fetchall()

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Name", "Email", "Phone", "Department", "Year", "Registration Number", "Event", "Competition"])
    for row in rows:
        writer.writerow([row["name"], row["email"], row["phone"], row["department"], row["year"], row["registration_no"], row["event_name"], row["competition_name"]])

    response = StreamingResponse(iter([output.getvalue()]), media_type="text/csv")
    response.headers["Content-Disposition"] = "attachment; filename=participants_report.csv"
    return response


@app.get("/api/reports/participants/xlsx")
async def export_participants_xlsx(event_id: Optional[int] = Query(None), competition_id: Optional[int] = Query(None)):
    with get_connection() as conn:
        if competition_id is not None:
            rows = conn.execute(
                "SELECT p.name, p.email, p.phone, p.department, p.year, p.registration_no, e.name AS event_name, c.name AS competition_name FROM participants p JOIN events e ON e.id = p.event_id JOIN competitions c ON c.id = p.competition_id WHERE p.competition_id = ? ORDER BY p.created_at DESC",
                (competition_id,),
            ).fetchall()
        elif event_id is not None:
            rows = conn.execute(
                "SELECT p.name, p.email, p.phone, p.department, p.year, p.registration_no, e.name AS event_name, c.name AS competition_name FROM participants p JOIN events e ON e.id = p.event_id JOIN competitions c ON c.id = p.competition_id WHERE p.event_id = ? ORDER BY p.created_at DESC",
                (event_id,),
            ).fetchall()
        else:
            rows = conn.execute(
                "SELECT p.name, p.email, p.phone, p.department, p.year, p.registration_no, e.name AS event_name, c.name AS competition_name FROM participants p JOIN events e ON e.id = p.event_id JOIN competitions c ON c.id = p.competition_id ORDER BY p.created_at DESC"
            ).fetchall()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Participants"
    headers = ["Name", "Email", "Phone", "Department", "Year", "Registration Number", "Event", "Competition"]
    sheet.append(headers)
    for row in rows:
        sheet.append([
            row["name"],
            row["email"],
            row["phone"],
            row["department"],
            row["year"],
            row["registration_no"],
            row["event_name"],
            row["competition_name"],
        ])

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)

    response = StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response.headers["Content-Disposition"] = "attachment; filename=participants_report.xlsx"
    return response


@app.get("/api/reports/volunteers/csv")
async def export_volunteers_csv(event_id: Optional[int] = Query(None), competition_id: Optional[int] = Query(None)):
    with get_connection() as conn:
        if competition_id is not None:
            rows = conn.execute(
                "SELECT v.name, v.email, v.phone, v.preferred_role, v.skills, v.assignment, e.name AS event_name, c.name AS competition_name FROM volunteers v JOIN events e ON e.id = v.event_id JOIN competitions c ON c.id = v.competition_id WHERE v.competition_id = ? ORDER BY v.created_at DESC",
                (competition_id,),
            ).fetchall()
        elif event_id is not None:
            rows = conn.execute(
                "SELECT v.name, v.email, v.phone, v.preferred_role, v.skills, v.assignment, e.name AS event_name, c.name AS competition_name FROM volunteers v JOIN events e ON e.id = v.event_id JOIN competitions c ON c.id = v.competition_id WHERE v.event_id = ? ORDER BY v.created_at DESC",
                (event_id,),
            ).fetchall()
        else:
            rows = conn.execute(
                "SELECT v.name, v.email, v.phone, v.preferred_role, v.skills, v.assignment, e.name AS event_name, c.name AS competition_name FROM volunteers v JOIN events e ON e.id = v.event_id JOIN competitions c ON c.id = v.competition_id ORDER BY v.created_at DESC"
            ).fetchall()

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Name", "Email", "Phone", "Preferred Role", "Skills", "Assignment", "Event", "Competition"])
    for row in rows:
        writer.writerow([row["name"], row["email"], row["phone"], row["preferred_role"], row["skills"], row["assignment"], row["event_name"], row["competition_name"]])

    response = StreamingResponse(iter([output.getvalue()]), media_type="text/csv")
    response.headers["Content-Disposition"] = "attachment; filename=volunteers_report.csv"
    return response


@app.get("/api/reports/volunteers/xlsx")
async def export_volunteers_xlsx(event_id: Optional[int] = Query(None), competition_id: Optional[int] = Query(None)):
    with get_connection() as conn:
        if competition_id is not None:
            rows = conn.execute(
                "SELECT v.name, v.email, v.phone, v.preferred_role, v.skills, v.assignment, e.name AS event_name, c.name AS competition_name FROM volunteers v JOIN events e ON e.id = v.event_id JOIN competitions c ON c.id = v.competition_id WHERE v.competition_id = ? ORDER BY v.created_at DESC",
                (competition_id,),
            ).fetchall()
        elif event_id is not None:
            rows = conn.execute(
                "SELECT v.name, v.email, v.phone, v.preferred_role, v.skills, v.assignment, e.name AS event_name, c.name AS competition_name FROM volunteers v JOIN events e ON e.id = v.event_id JOIN competitions c ON c.id = v.competition_id WHERE v.event_id = ? ORDER BY v.created_at DESC",
                (event_id,),
            ).fetchall()
        else:
            rows = conn.execute(
                "SELECT v.name, v.email, v.phone, v.preferred_role, v.skills, v.assignment, e.name AS event_name, c.name AS competition_name FROM volunteers v JOIN events e ON e.id = v.event_id JOIN competitions c ON c.id = v.competition_id ORDER BY v.created_at DESC"
            ).fetchall()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Volunteers"
    headers = ["Name", "Email", "Phone", "Preferred Role", "Skills", "Assignment", "Event", "Competition"]
    sheet.append(headers)
    for row in rows:
        sheet.append([
            row["name"],
            row["email"],
            row["phone"],
            row["preferred_role"],
            row["skills"],
            row["assignment"],
            row["event_name"],
            row["competition_name"],
        ])

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)

    response = StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response.headers["Content-Disposition"] = "attachment; filename=volunteers_report.xlsx"
    return response


@app.get("/api/reports/combined/xlsx")
async def export_combined_xlsx(event_id: Optional[int] = Query(None), competition_id: Optional[int] = Query(None)):
    with get_connection() as conn:
        if competition_id is not None:
            participant_rows = conn.execute(
                "SELECT p.name, p.email, p.phone, p.department, p.year, p.registration_no, e.name AS event_name, c.name AS competition_name FROM participants p JOIN events e ON e.id = p.event_id JOIN competitions c ON c.id = p.competition_id WHERE p.competition_id = ? ORDER BY p.created_at DESC",
                (competition_id,),
            ).fetchall()
            volunteer_rows = conn.execute(
                "SELECT v.name, v.email, v.phone, v.preferred_role, v.skills, v.assignment, e.name AS event_name, c.name AS competition_name FROM volunteers v JOIN events e ON e.id = v.event_id JOIN competitions c ON c.id = v.competition_id WHERE v.competition_id = ? ORDER BY v.created_at DESC",
                (competition_id,),
            ).fetchall()
        elif event_id is not None:
            participant_rows = conn.execute(
                "SELECT p.name, p.email, p.phone, p.department, p.year, p.registration_no, e.name AS event_name, c.name AS competition_name FROM participants p JOIN events e ON e.id = p.event_id JOIN competitions c ON c.id = p.competition_id WHERE p.event_id = ? ORDER BY p.created_at DESC",
                (event_id,),
            ).fetchall()
            volunteer_rows = conn.execute(
                "SELECT v.name, v.email, v.phone, v.preferred_role, v.skills, v.assignment, e.name AS event_name, c.name AS competition_name FROM volunteers v JOIN events e ON e.id = v.event_id JOIN competitions c ON c.id = v.competition_id WHERE v.event_id = ? ORDER BY v.created_at DESC",
                (event_id,),
            ).fetchall()
        else:
            participant_rows = conn.execute(
                "SELECT p.name, p.email, p.phone, p.department, p.year, p.registration_no, e.name AS event_name, c.name AS competition_name FROM participants p JOIN events e ON e.id = p.event_id JOIN competitions c ON c.id = p.competition_id ORDER BY p.created_at DESC"
            ).fetchall()
            volunteer_rows = conn.execute(
                "SELECT v.name, v.email, v.phone, v.preferred_role, v.skills, v.assignment, e.name AS event_name, c.name AS competition_name FROM volunteers v JOIN events e ON e.id = v.event_id JOIN competitions c ON c.id = v.competition_id ORDER BY v.created_at DESC"
            ).fetchall()

    workbook = Workbook()
    participant_sheet = workbook.active
    participant_sheet.title = "Participants"
    participant_sheet.append(["Name", "Email", "Phone", "Department", "Year", "Registration Number", "Event", "Competition"])
    for row in participant_rows:
        participant_sheet.append([row["name"], row["email"], row["phone"], row["department"], row["year"], row["registration_no"], row["event_name"], row["competition_name"]])

    volunteer_sheet = workbook.create_sheet(title="Volunteers")
    volunteer_sheet.append(["Name", "Email", "Phone", "Preferred Role", "Skills", "Assignment", "Event", "Competition"])
    for row in volunteer_rows:
        volunteer_sheet.append([row["name"], row["email"], row["phone"], row["preferred_role"], row["skills"], row["assignment"], row["event_name"], row["competition_name"]])

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)

    response = StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response.headers["Content-Disposition"] = "attachment; filename=combined_report.xlsx"
    return response


@app.post("/api/participants")
async def create_participant(payload: ParticipantCreate):
    with get_connection() as conn:
        event = conn.execute("SELECT id FROM events WHERE id = ?", (payload.event_id,)).fetchone()
        if not event:
            raise HTTPException(status_code=404, detail="Selected event does not exist.")
        competition = conn.execute(
            "SELECT id, capacity FROM competitions WHERE id = ? AND event_id = ?",
            (payload.competition_id, payload.event_id),
        ).fetchone()
        if not competition:
            raise HTTPException(status_code=404, detail="Selected competition does not exist in the event.")

        duplicate = conn.execute(
            "SELECT id FROM participants WHERE competition_id = ? AND (email = ? OR phone = ?)",
            (payload.competition_id, payload.email.strip().lower(), payload.phone.strip()),
        ).fetchone()
        if duplicate:
            raise HTTPException(status_code=400, detail="This participant is already registered for this competition.")

        current_count = conn.execute(
            "SELECT COUNT(*) AS total FROM participants WHERE competition_id = ?",
            (payload.competition_id,),
        ).fetchone()["total"]
        if current_count >= competition["capacity"]:
            raise HTTPException(status_code=400, detail="Competition capacity is full.")

        role_conflict = conn.execute(
            "SELECT id FROM volunteers WHERE competition_id = ? AND event_id = ? AND (email = ? OR phone = ?)",
            (payload.competition_id, payload.event_id, payload.email.strip().lower(), payload.phone.strip()),
        ).fetchone()
        if role_conflict:
            raise HTTPException(status_code=400, detail="This person is already assigned as a volunteer for this competition.")

        cursor = conn.execute(
            """
            INSERT INTO participants (event_id, competition_id, name, email, phone, department, year, registration_no)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                payload.event_id,
                payload.competition_id,
                payload.name.strip(),
                payload.email.strip().lower(),
                payload.phone.strip(),
                payload.department.strip(),
                payload.year.strip(),
                payload.registration_no.strip(),
            ),
        )
        conn.commit()
        participant_id = cursor.lastrowid
    return {"id": participant_id, "message": "Participant registered successfully."}


@app.patch("/api/participants/{participant_id}")
async def update_participant(participant_id: int, payload: dict):
    with get_connection() as conn:
        existing = conn.execute("SELECT * FROM participants WHERE id = ?", (participant_id,)).fetchone()
        if not existing:
            raise HTTPException(status_code=404, detail="Participant not found.")

        updates = {}
        for key in ["name", "email", "phone", "department", "year", "registration_no"]:
            if key in payload and payload[key] is not None:
                updates[key] = str(payload[key]).strip()
        if not updates:
            raise HTTPException(status_code=400, detail="No valid update fields provided.")

        fields = []
        values = []
        for key, value in updates.items():
            fields.append(f"{key} = ?")
            values.append(value)
        values.append(participant_id)

        conn.execute(f"UPDATE participants SET {', '.join(fields)} WHERE id = ?", tuple(values))
        conn.commit()
        row = conn.execute("SELECT * FROM participants WHERE id = ?", (participant_id,)).fetchone()
    return dict(row)


@app.delete("/api/participants/{participant_id}")
async def delete_participant(participant_id: int):
    with get_connection() as conn:
        cursor = conn.execute("DELETE FROM participants WHERE id = ?", (participant_id,))
        conn.commit()
    if cursor.rowcount == 0:
        raise HTTPException(status_code=404, detail="Participant not found.")
    return {"message": "Participant deleted successfully."}


@app.get("/api/volunteers")
async def list_volunteers(competition_id: Optional[int] = Query(None)):
    with get_connection() as conn:
        if competition_id is not None:
            rows = conn.execute(
                "SELECT * FROM volunteers WHERE competition_id = ? ORDER BY created_at DESC",
                (competition_id,),
            ).fetchall()
        else:
            rows = conn.execute(
                "SELECT * FROM volunteers ORDER BY created_at DESC"
            ).fetchall()
    return [dict(row) for row in rows]


@app.post("/api/volunteers")
async def create_volunteer(payload: VolunteerCreate):
    with get_connection() as conn:
        event = conn.execute("SELECT id FROM events WHERE id = ?", (payload.event_id,)).fetchone()
        if not event:
            raise HTTPException(status_code=404, detail="Selected event does not exist.")
        competition = conn.execute(
            "SELECT id FROM competitions WHERE id = ? AND event_id = ?",
            (payload.competition_id, payload.event_id),
        ).fetchone()
        if not competition:
            raise HTTPException(status_code=404, detail="Selected competition does not exist in the event.")

        duplicate = conn.execute(
            "SELECT id FROM volunteers WHERE competition_id = ? AND (email = ? OR phone = ?)",
            (payload.competition_id, payload.email.strip().lower(), payload.phone.strip()),
        ).fetchone()
        if duplicate:
            raise HTTPException(status_code=400, detail="This volunteer application already exists for this competition.")

        participant_conflict = conn.execute(
            "SELECT id FROM participants WHERE competition_id = ? AND event_id = ? AND (email = ? OR phone = ?)",
            (payload.competition_id, payload.event_id, payload.email.strip().lower(), payload.phone.strip()),
        ).fetchone()
        if participant_conflict:
            raise HTTPException(status_code=400, detail="This person is already registered as a participant for this competition.")

        cursor = conn.execute(
            """
            INSERT INTO volunteers (event_id, competition_id, name, email, phone, preferred_role, skills, assignment)
            VALUES (?, ?, ?, ?, ?, ?, ?, 'Pending')
            """,
            (
                payload.event_id,
                payload.competition_id,
                payload.name.strip(),
                payload.email.strip().lower(),
                payload.phone.strip(),
                payload.preferred_role.strip(),
                payload.skills.strip(),
            ),
        )
        conn.commit()
        volunteer_id = cursor.lastrowid
    return {"id": volunteer_id, "message": "Volunteer request submitted successfully."}


@app.patch("/api/volunteers/{volunteer_id}")
async def update_volunteer(volunteer_id: int, payload: dict):
    with get_connection() as conn:
        existing = conn.execute("SELECT * FROM volunteers WHERE id = ?", (volunteer_id,)).fetchone()
        if not existing:
            raise HTTPException(status_code=404, detail="Volunteer not found.")

        updates = {}
        for key in ["name", "email", "phone", "preferred_role", "skills", "assignment"]:
            if key in payload and payload[key] is not None:
                updates[key] = str(payload[key]).strip()
        if not updates:
            raise HTTPException(status_code=400, detail="No valid update fields provided.")

        fields = []
        values = []
        for key, value in updates.items():
            fields.append(f"{key} = ?")
            values.append(value)
        values.append(volunteer_id)

        conn.execute(f"UPDATE volunteers SET {', '.join(fields)} WHERE id = ?", tuple(values))
        conn.commit()
        row = conn.execute("SELECT * FROM volunteers WHERE id = ?", (volunteer_id,)).fetchone()
    return dict(row)


@app.delete("/api/volunteers/{volunteer_id}")
async def delete_volunteer(volunteer_id: int):
    with get_connection() as conn:
        cursor = conn.execute("DELETE FROM volunteers WHERE id = ?", (volunteer_id,))
        conn.commit()
    if cursor.rowcount == 0:
        raise HTTPException(status_code=404, detail="Volunteer not found.")
    return {"message": "Volunteer deleted successfully."}


@app.get("/api/dashboard")
async def dashboard_summary():
    with get_connection() as conn:
        event_count = conn.execute("SELECT COUNT(*) AS total FROM events").fetchone()["total"]
        competition_count = conn.execute("SELECT COUNT(*) AS total FROM competitions").fetchone()["total"]
        participant_count = conn.execute("SELECT COUNT(*) AS total FROM participants").fetchone()["total"]
        volunteer_count = conn.execute("SELECT COUNT(*) AS total FROM volunteers").fetchone()["total"]

        recent_participants = conn.execute(
            """
            SELECT p.id, p.name, e.name AS event_name, c.name AS competition_name, p.created_at
            FROM participants p
            JOIN events e ON e.id = p.event_id
            JOIN competitions c ON c.id = p.competition_id
            ORDER BY p.created_at DESC
            LIMIT 5
            """
        ).fetchall()

        recent_volunteers = conn.execute(
            """
            SELECT v.id, v.name, e.name AS event_name, c.name AS competition_name, v.assignment
            FROM volunteers v
            JOIN events e ON e.id = v.event_id
            JOIN competitions c ON c.id = v.competition_id
            ORDER BY v.created_at DESC
            LIMIT 5
            """
        ).fetchall()

        all_participants = conn.execute(
            """
            SELECT p.id, p.name, p.email, p.phone, e.name AS event_name, c.name AS competition_name
            FROM participants p
            JOIN events e ON e.id = p.event_id
            JOIN competitions c ON c.id = p.competition_id
            ORDER BY p.created_at DESC
            """
        ).fetchall()

        all_volunteers = conn.execute(
            """
            SELECT v.id, v.name, v.email, v.phone, v.preferred_role, v.assignment, e.name AS event_name, c.name AS competition_name
            FROM volunteers v
            JOIN events e ON e.id = v.event_id
            JOIN competitions c ON c.id = v.competition_id
            ORDER BY v.created_at DESC
            """
        ).fetchall()

        event_breakdown = conn.execute(
            """
            SELECT
                e.id,
                e.name,
                (SELECT COUNT(*) FROM competitions c WHERE c.event_id = e.id) AS competitions,
                (SELECT COUNT(*) FROM participants p WHERE p.event_id = e.id) AS participants,
                (SELECT COUNT(*) FROM volunteers v WHERE v.event_id = e.id) AS volunteers
            FROM events e
            ORDER BY e.date DESC, e.id DESC
            """
        ).fetchall()

    return {
        "stats": {
            "events": event_count,
            "competitions": competition_count,
            "participants": participant_count,
            "volunteers": volunteer_count,
        },
        "recent_participants": [dict(row) for row in recent_participants],
        "recent_volunteers": [dict(row) for row in recent_volunteers],
        "all_participants": [dict(row) for row in all_participants],
        "all_volunteers": [dict(row) for row in all_volunteers],
        "event_breakdown": [dict(row) for row in event_breakdown],
    }


if __name__ == "__main__":
    import uvicorn

    uvicorn.run("app:app", host="0.0.0.0", port=8000, reload=True)
